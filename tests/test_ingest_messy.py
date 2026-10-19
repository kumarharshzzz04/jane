"""Regression tests: title/metadata rows above the header, ragged CSVs, merged-cell Excel."""
import os
import tempfile

import pandas as pd
import pytest

from jane.ingest import _detect_delimiter, _find_header_row, _split_rows, load_file


def _load(content: bytes) -> pd.DataFrame:
    fd, tmp = tempfile.mkstemp(suffix=".csv")
    p = None
    try:
        p = __import__("pathlib").Path(tmp)
        p.write_bytes(content)
        return load_file(p).primary
    finally:
        os.close(fd)
        if p is not None:
            try:
                p.unlink()
            except PermissionError:
                pass


CASES = {
    "title_single": b"Monthly Export Report\nregion,amount\nnorth,10\nsouth,20\n",
    "title_two_lines": b"Monthly Export Report\nPrepared by ACME\nregion,amount,note\nnorth,10,ok\nsouth,20,ok\n",
    "metadata_below": b"source: CRM export\nregion,amount,note\nnorth,10,ok\nsouth,20,ok\n",
    "ragged_body": b"Monthly Export Report\nregion,amount,note\nnorth,10\nsouth,20,ok\n",
}
EXPECTED = {
    "title_single": ["region", "amount"],
    "title_two_lines": ["region", "amount", "note"],
    "metadata_below": ["region", "amount", "note"],
    "ragged_body": ["region", "amount", "note"],
}


@pytest.mark.parametrize("name", sorted(CASES))
def test_title_rows_are_skipped(name):
    df = _load(CASES[name])
    assert list(df.columns) == EXPECTED[name]
    assert "Monthly" not in " ".join(map(str, df.columns))


def test_leading_junk_inside_table_is_preserved():
    # junk row *below* a valid header must NOT trigger a reslice
    df = _load(b"region,amount\n,Monthly Export Report\nnorth,10\nsouth,20\n")
    assert list(df.columns) == ["region", "amount"]
    assert len(df) == 3


def test_detect_delimiter_prefers_consistent_commas():
    assert _detect_delimiter(b"a,b,c\n1,2,3\nx,y,z\n") == ","
    assert _detect_delimiter(b"a\tb\n1\t2\nx\ty\n") == "\t"


def test_find_header_row_title():
    rows = _split_rows(b"Hello World\na,b\n1,2\n", "utf-8", ",", 10)
    assert _find_header_row(rows) == 1


def test_find_header_row_clean():
    rows = _split_rows(b"a,b\n1,2\n3,4\n", "utf-8", ",", 10)
    assert _find_header_row(rows) == 0


def test_merged_cell_excel(tmp_path):
    openpyxl = pytest.importorskip("openpyxl")
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "rep"
    ws["A2"], ws["B2"], ws["C2"] = "region", "amount", "note"
    ws["A3"], ws["B3"], ws["C3"] = "north", 10, "ok"
    ws["A4"], ws["C4"] = "south", "ok"  # B4 left empty; merged-style gap
    ws.merge_cells("A5:A6")
    ws["A5"] = "east"
    ws["B5"], ws["C5"] = 30, "ok"
    ws["B6"], ws["C6"] = 40, "ok"
    p = tmp_path / "messy.xlsx"
    wb.save(p)

    loaded = load_file(p)
    df = loaded.sheets[0].df
    assert list(df.columns) == ["region", "amount", "note"]
    # merged A5:A6 paints 'east' into both rows
    assert df[df["region"] == "east"].shape[0] == 2
